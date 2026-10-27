"""Isolated app + database per test, with the LLM stubbed out."""

import io
import os
import tempfile
from collections.abc import Iterator
from typing import Any

import docx
import pytest
from fastapi.testclient import TestClient

# Point the whole app at a throwaway data dir BEFORE anything imports app.*,
# so no test can ever open the developer's real database or config.
os.environ["DATA_DIR"] = tempfile.mkdtemp(prefix="lokishift-tests-")

JD = (
    "Senior Python Developer\n\nRequirements:\n- 5+ years of Python\n- FastAPI, Django\n"
    "- Kafka, Kubernetes, Terraform\n- AWS\n\nResponsibilities:\n- Build REST APIs\n"
    "- Run event-driven microservices\n"
)


def make_docx(bullets: list[str] | None = None) -> bytes:
    document = docx.Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Email: jane@example.com | Phone: 512-555-0100")
    document.add_paragraph("Professional Summary:")
    document.add_paragraph("Backend engineer with 8 years of Python experience building APIs.")
    document.add_paragraph("Skills: Python, Django, PostgreSQL, Docker")
    document.add_paragraph("Work Experience:")
    document.add_paragraph("Client: Acme Bank, Austin, TX")
    document.add_paragraph("Duration: Jan 2021 - Present")
    for text in bullets or [
        "Built Django REST APIs for account services.",
        "Wrote data pipelines in Python for nightly reports.",
    ]:
        document.add_paragraph(text, style="List Bullet")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def stub_llm(monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    """Replace the tiered LLM call. Ranking prefers resumes mentioning Kafka."""
    calls: dict[str, Any] = {"rank": 0, "tailor": 0, "fail_tailor": None}

    async def fake(task: str, **kwargs: Any) -> tuple[dict[str, Any], str]:
        from app.llm import get_llm_config
        from app.services.matching import ProviderBillingError

        model = get_llm_config().model
        calls.setdefault("models", []).append(model)
        if model in calls.get("broken_models", set()):
            raise ProviderBillingError()
        prompt = kwargs["prompt"]
        if task == "rank":
            calls["rank"] += 1
            blocks = prompt.split('=== base_id: "')[1:]
            rankings = []
            for block in blocks:
                base_id = block.split('"', 1)[0]
                rankings.append({"base_id": base_id, "score": 90 if "Kafka" in block else 60,
                                 "strengths": ["Python"], "gaps": [], "reasoning": "stub"})
            rankings.sort(key=lambda r: -r["score"])
            return {"company": "", "role": "Senior Python Developer", "rankings": rankings,
                    "recommended_base_id": rankings[0]["base_id"]}, "stub"
        calls["tailor"] += 1
        if calls["fail_tailor"]:
            raise calls["fail_tailor"]
        lines = [l for l in prompt.splitlines() if l.startswith("[") and "LOCKED" not in l]
        indexes = {int(l[1:l.index("]")]): l for l in lines}
        summary = next(i for i, l in indexes.items() if "Backend engineer" in l)
        bullet = next(i for i, l in indexes.items() if "(List Bullet)" in l)
        phone = 1  # protected: must be ignored even if the model returns it
        return {
            "summary": "stub",
            "edits": [
                {"index": summary, "new_text": "Senior Python developer with 8 years building FastAPI and Django APIs on AWS.", "keywords": ["FastAPI"]},
                {"index": phone, "new_text": "Phone: 000", "keywords": []},
            ],
            "additions": [
                {"after": bullet, "text": "Built event-driven microservices with Kafka on Kubernetes.", "keywords": ["Kafka"]},
            ],
        }, "stub"

    import app.services.matching as matching
    import app.services.tailoring as tailoring

    monkeypatch.setattr(matching, "_complete_json_tiered", fake)
    monkeypatch.setattr(tailoring, "_complete_json_tiered", fake)
    return calls


@pytest.fixture
def configure_llm() -> bool:
    return True


@pytest.fixture
def client(
    tmp_path, monkeypatch: pytest.MonkeyPatch, stub_llm, configure_llm: bool
) -> Iterator[TestClient]:
    from app import database
    from app.config import settings
    from app.routers import connections, health, profiles, tasks
    from app.services import worker

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    test_db = database.Database(tmp_path / "test.db")
    import app.connections as conn_service

    for module in (database, profiles, tasks, worker, health, connections, conn_service):
        if hasattr(module, "db"):
            monkeypatch.setattr(module, "db", test_db)
    import app.services.matching as matching

    monkeypatch.setattr(matching, "db", test_db)
    monkeypatch.setattr(settings, "job_rate_limit_retries", 0)
    from app.main import app

    with TestClient(app) as c:
        if configure_llm:
            # The only supported way to configure the AI: a connection on the Settings page.
            r = c.post(
                "/api/v1/connections",
                json={"name": "OpenAI", "provider": "openai", "model": "gpt-test", "api_key": "sk-test-settings"},
            )
            assert r.status_code == 200, r.text
        yield c
